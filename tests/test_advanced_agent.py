from __future__ import annotations

import io
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from backend.agent.contracts import ModelResponse, ModelToolCall, TaskContract, ToolSpec
from backend.agent.context import ContextBuilder
from backend.agent.loop import AgentLoop, _should_wrap_up, _tool_message
from backend.agent.model import (
    ChatCompletionsAdapter,
    ModelProtocolError,
    ResponsesAdapter,
    ScriptedModelAdapter,
    build_model_adapter,
)
from backend.agent.store import RunStore
from backend.agent.tools import ToolExecutor, ToolRegistry, _validate
from backend.services.advanced_agent import FORMAL_AGENT_TOOLS, _run_knowledge_ids
from backend.services.data_plane.contracts import DatasetRef, DatasetRefStore
from backend.services.data_plane.sandbox import SandboxUnavailable
from backend.services.data_plane.trino import TrinoAdapter, TrinoConfig
from backend.services.results.manifests import ResultService
from backend.services.skills import get_skill, unsupported_formal_tools


def _contract(source_ids=(), objective="按区域核对销售额"):
    return TaskContract.from_payload({
        "objective": objective, "coverage": "所选数据的全部记录",
        "dimensions": ["区域"], "deliverables": ["summary"],
        "source_scope": list(source_ids),
    })


def _confirmed_run(
    app, *, source_ids=(), allowed=("query", "validate"), budget=None,
    objective="按区域核对销售额",
):
    database = app.extensions["meridian_db"]
    session = database.put(
        "sessions", {"id": database.new_id("ses"), "workspace_id": "default", "owner_id": "local-default"},
        workspace_id="default",
    )
    store = RunStore(database)
    run, _ = store.create_run(
        workspace_id="default", session_id=session["id"], actor_id="local-default",
        source_scope=list(source_ids), allowed_tool_ids=list(allowed), budget=budget,
    )
    store.add_contract(
        run["id"], _contract(source_ids, objective),
        expected_version=0, confirmed_by="local-default",
    )
    store.add_plan(run["id"], {
        "tasks": [{"id": "analyze", "title": "分析", "status": "open", "depends_on": []}],
    }, reason="test", expected_version=0)
    return store, store.get_run(run["id"])


def test_analysis_api_requires_versioned_confirmation_and_uses_typed_job(client, source):
    created = client.post("/api/analyses", json={
        "objective": "分析区域销售", "source_ids": [source["id"]],
    })
    assert created.status_code == 201
    run = created.get_json()["item"]
    assert run["execution_status"] == "waiting_input"
    assert run["contract"]["version"] == 1
    assert not run["contract"]["confirmed_at"]

    conflict = client.post(
        f"/api/analyses/{run['id']}/contract/confirm", json={"expected_version": 0},
    )
    assert conflict.status_code == 400
    confirmed = client.post(
        f"/api/analyses/{run['id']}/contract/confirm", json={"expected_version": 1},
    )
    assert confirmed.status_code == 200
    job = confirmed.get_json()["job"]
    assert job["typed"] is True


def test_formal_analysis_exposes_only_supported_skills(client, source):
    with client.application.app_context():
        assert unsupported_formal_tools(get_skill("quality-audit", "default")) == []
    listed = client.get("/api/skills").get_json()["items"]
    assert {item["id"] for item in listed} == {"executive-summary", "quality-audit", "trend-diagnosis"}
    assert all(item["formal_compatible"] for item in listed)
    rejected = client.post("/api/analyses", json={
        "objective": "评估业务", "source_ids": [source["id"]], "skill_id": "bcg-matrix",
    })
    assert rejected.status_code == 400
    agent = client.post("/api/agents", json={
        "name": "不兼容智能体", "source_ids": [source["id"]], "skill_id": "bcg-matrix",
    })
    assert agent.status_code == 400


@pytest.mark.parametrize("mode,expected", [
    ("quick", 100_000), ("auto", 160_000), ("deep", 240_000),
])
def test_interactive_analysis_budget_matches_execution_mode(client, mode, expected):
    response = client.post("/api/analyses", json={
        "objective": "分析成绩分布", "execution_mode": mode,
        "confirm_required": True,
    })
    assert response.status_code == 201
    assert response.get_json()["item"]["budget"]["model_tokens"] == expected


def test_analysis_attachments_are_bounded_and_contract_scope_locks(client):
    run = client.post("/api/analyses", json={"objective": "读取业务口径"}).get_json()["item"]
    uploaded = client.post(
        f"/api/analyses/{run['id']}/attachments",
        data={"files": (io.BytesIO("指标定义：销售额。".encode()), "definition.md")},
        content_type="multipart/form-data",
    )
    assert uploaded.status_code == 201
    item = uploaded.get_json()["items"][0]
    assert item["format"] == "md"
    rejected = client.post(
        f"/api/analyses/{run['id']}/attachments",
        data={"files": (io.BytesIO(b"bad"), "script.py")}, content_type="multipart/form-data",
    )
    assert rejected.status_code == 400
    assert client.post(
        f"/api/analyses/{run['id']}/contract/confirm", json={"expected_version": 1},
    ).status_code == 200
    assert client.delete(f"/api/analyses/{run['id']}/attachments/{item['id']}").status_code == 400


def test_analysis_attachment_stays_private_and_follows_its_run_lineage(app, client):
    parent = client.post("/api/analyses", json={"objective": "核对私有口径"}).get_json()["item"]
    response = client.post(
        f"/api/analyses/{parent['id']}/attachments",
        data={"files": (io.BytesIO("私有口径：核算已授权订单。".encode()), "private.md")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 201
    document_id = response.get_json()["items"][0]["document_id"]
    assert document_id not in {
        item["id"] for item in client.get("/api/knowledge/documents").get_json()["items"]
    }
    assert client.post("/api/knowledge/search", json={"query": "私有口径"}).get_json()["items"] == []
    assert client.post("/api/analyses", json={
        "objective": "错误引用", "knowledge_document_ids": [document_id],
    }).status_code == 400

    unrelated = client.post("/api/analyses", json={"objective": "另一项分析"}).get_json()["item"]
    child_response = client.post(
        f"/api/analyses/{parent['id']}/branch", json={"mode": "followup", "prompt": "继续核对口径"},
    )
    assert child_response.status_code == 201
    child = child_response.get_json()["item"]
    database = app.extensions["meridian_db"]
    with app.app_context():
        assert document_id in _run_knowledge_ids(database, parent)
        assert document_id in _run_knowledge_ids(database, child)
        assert document_id not in _run_knowledge_ids(database, unrelated)


def test_attachment_batch_failure_does_not_leave_partial_documents(client):
    run = client.post("/api/analyses", json={"objective": "读取两份定义"}).get_json()["item"]
    response = client.post(
        f"/api/analyses/{run['id']}/attachments",
        data={"files": [
            (io.BytesIO("有效定义".encode()), "valid.md"),
            (io.BytesIO(b""), "empty.md"),
        ]},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert client.get(f"/api/analyses/{run['id']}/attachments").get_json()["items"] == []
    assert client.get("/api/knowledge/documents").get_json()["items"] == []


def test_single_agent_loop_publishes_only_after_independent_validation(app, client):
    store, run = _confirmed_run(app)
    store.db.put("query_results", {
        "id": "result-1", "workspace_id": "default", "source_ids": [],
        "rows": 1, "columns": ["checked"], "data": [{"checked": 1}],
        "completeness": "complete", "accuracy": "exact",
    }, workspace_id="default")
    registry = ToolRegistry()
    registry.register(ToolSpec("query", "query", {"type": "object", "properties": {}}), lambda _args: {
        "result_id": "result-1", "output_refs": ["result-1"], "completeness": "complete",
    })
    registry.register(ToolSpec("validate", "validate", {"type": "object", "properties": {}}), lambda _args: {
        "output_refs": ["result-1"], "completeness": "complete", "validation_status": "PASS",
    })
    model = ScriptedModelAdapter([
        ModelResponse("scripted_test", "fixture", "", (
            ModelToolCall("q1", "query", {}), ModelToolCall("v1", "validate", {}),
        ), "tool_calls", None, {"total_tokens": 20}),
        ModelResponse("scripted_test", "fixture", "销售额已核对。", (), "stop", None, {"total_tokens": 5}),
    ])
    loop = AgentLoop(
        store=store, model=model, tools=ToolExecutor(store, registry),
        finalizer=ResultService(store.db).finalize,
    )
    result = loop.run(run["id"], runner_id="test-runner", history=[])
    assert result.status == "finished"
    assert result.outcome == "complete"
    assert result.publication_id
    assert {item["tool_id"] for item in store.actions(run["id"])} == {"query", "validate"}
    feedback = client.post("/api/feedback", json={
        "run_id": run["id"], "rating": "incorrect", "comment": "请复核统计口径",
    })
    assert feedback.status_code == 201
    item = feedback.get_json()["item"]
    assert item["status"] == "open"
    assert client.get(f"/api/analyses/{run['id']}/feedback").get_json()["item"]["rating"] == "incorrect"
    changed = client.post("/api/feedback", json={
        "run_id": run["id"], "rating": "partially_correct", "comment": "有一项口径需复查",
    })
    assert changed.status_code == 200
    assert changed.get_json()["item"]["id"] == item["id"]
    reviewed = client.patch(f"/api/feedback/{item['id']}", json={
        "status": "resolved", "review_note": "已核对口径",
    })
    assert reviewed.status_code == 200
    assert reviewed.get_json()["item"]["status"] == "resolved"
    artifacts = client.post(f"/api/analyses/{run['id']}/artifacts", json={
        "kinds": ["summary_docx", "report_docx", "dashboard_png"],
    })
    assert artifacts.status_code == 201
    assert {item["kind"] for item in artifacts.get_json()["items"]} == {
        "summary_docx", "report_docx", "dashboard_png",
    }
    result_path = app.config["SETTINGS"].export_dir / "result-1.csv"
    result_path.write_text('checked,=HYPERLINK("https://example.test")\n1,=SUM(1+1)\n', encoding="utf-8")
    store.db.patch("query_results", "result-1", {"path": str(result_path)}, workspace_id="default")
    more_artifacts = client.post(f"/api/analyses/{run['id']}/artifacts", json={
        "kinds": ["data_xlsx", "report_pptx"],
    })
    assert more_artifacts.status_code == 201, more_artifacts.get_json()
    from openpyxl import load_workbook
    from pptx import Presentation

    created = {item["kind"]: item for item in more_artifacts.get_json()["items"]}
    workbook = load_workbook(store.db.get("artifacts", created["data_xlsx"]["id"])["path"], read_only=True)
    assert workbook["验证数据"]["A2"].value == 1
    assert workbook["验证数据"]["B1"].data_type != "f"
    assert workbook["验证数据"]["B1"].value.startswith("'=")
    assert workbook["验证数据"]["B2"].data_type != "f"
    assert workbook["验证数据"]["B2"].value.startswith("'=")
    presentation = Presentation(store.db.get("artifacts", created["report_pptx"]["id"])["path"])
    assert len(presentation.slides) >= 3
    email = client.post(f"/api/analyses/{run['id']}/email/eml", json={
        "recipients": "reviewer@example.com", "kinds": ["summary_docx", "report_docx"],
    })
    assert email.status_code == 201
    eml = store.db.get("artifacts", email.get_json()["eml"]["id"], workspace_id="default")
    assert Path(eml["path"]).read_bytes().count(b"Content-Disposition: attachment") == 2


def test_agent_deterministically_removes_unverified_numeric_claim_before_retry(app):
    store, run = _confirmed_run(app)
    store.db.put("query_results", {
        "id": "result-repair", "workspace_id": "default", "source_ids": [],
        "rows": 1, "columns": ["total_sales"], "data": [{"total_sales": 745}],
        "completeness": "complete", "accuracy": "exact",
    }, workspace_id="default")
    registry = ToolRegistry()
    registry.register(ToolSpec("query", "query", {"type": "object", "properties": {}}), lambda _args: {
        "result_id": "result-repair", "output_refs": ["result-repair"], "completeness": "complete",
    })
    registry.register(ToolSpec("validate", "validate", {"type": "object", "properties": {}}), lambda _args: {
        "output_refs": ["result-repair"], "completeness": "complete", "validation_status": "PASS",
    })
    model = ScriptedModelAdapter([
        ModelResponse("scripted_test", "fixture", "", (
            ModelToolCall("q1", "query", {}), ModelToolCall("v1", "validate", {}),
        ), "tool_calls", None, {"total_tokens": 20}),
        ModelResponse(
            "scripted_test", "fixture",
            "销售额合计为 745 元。\n未经核对的预测为 999 元。",
            (), "stop", None, {"total_tokens": 5},
        ),
    ])
    result = AgentLoop(
        store=store, model=model, tools=ToolExecutor(store, registry),
        finalizer=ResultService(store.db).finalize,
    ).run(run["id"], runner_id="test-runner", history=[])
    assert result.outcome == "complete"
    assert result.answer == "销售额合计为 745 元。"
    assert len(store.decisions(run["id"])) == 2


def test_publication_gate_blocks_partial_or_unvalidated_result(app):
    store, run = _confirmed_run(app, allowed=("query",))
    registry = ToolRegistry()
    registry.register(ToolSpec("query", "query", {"type": "object", "properties": {}}), lambda _args: {
        "result_id": "partial-1", "output_refs": ["partial-1"], "completeness": "partial",
    })
    model = ScriptedModelAdapter([
        ModelResponse("scripted_test", "fixture", "", (ModelToolCall("q1", "query", {}),), "tool_calls", None, {"total_tokens": 1}),
        ModelResponse("scripted_test", "fixture", "不能作全量结论。", (), "stop", None, {"total_tokens": 1}),
    ])
    result = AgentLoop(
        store=store, model=model, tools=ToolExecutor(store, registry),
        finalizer=ResultService(store.db).finalize,
    ).run(run["id"], runner_id="test-runner", history=[])
    assert result.status == "finished"
    assert result.outcome == "partial"
    assert result.publication_id is None


def test_agent_strips_provider_think_blocks_before_finalization(app):
    store, run = _confirmed_run(app, allowed=())
    captured = {}

    def finalize(_run_id, answer, _evidence):
        captured["answer"] = answer
        return {"published": False, "quality_status": "failed"}

    model = ScriptedModelAdapter([ModelResponse(
        "scripted_test", "fixture", "<think>private reasoning 1 2 3</think>\n\n可见结论。",
        (), "stop", None, {"total_tokens": 1},
    )])
    result = AgentLoop(
        store=store, model=model, tools=ToolExecutor(store, ToolRegistry()), finalizer=finalize,
    ).run(run["id"], runner_id="test-runner", history=[])

    assert captured["answer"] == "可见结论。"
    assert result.answer == "可见结论。"


def test_agent_enforces_validation_before_publishing(app):
    store, run = _confirmed_run(app, allowed=("query_data", "validate_result"))
    registry = ToolRegistry()
    registry.register(
        ToolSpec("query_data", "query", {"type": "object", "properties": {}}),
        lambda _args: {"output_refs": ["result-1"], "completeness": "complete"},
    )
    registry.register(
        ToolSpec("validate_result", "validate", {"type": "object", "properties": {"result_id": {"type": "string"}}}),
        lambda args: {
            "output_refs": [args["result_id"]], "completeness": "complete",
            "validation_status": "PASS",
        },
    )
    model = ScriptedModelAdapter([
        ModelResponse(
            "scripted_test", "fixture", "", (ModelToolCall("q1", "query_data", {}),),
            "tool_calls", None, {"total_tokens": 1},
        ),
        ModelResponse("scripted_test", "fixture", "结果已核对。", (), "stop", None, {"total_tokens": 1}),
    ])

    def finalize(_run_id, answer, evidence):
        assert answer == "结果已核对。"
        assert any(
            item["tool"] == "validate_result" and item["validation_status"] == "PASS"
            for item in evidence
        )
        return {"published": True, "quality_status": "passed", "publication_id": "publication-1"}

    result = AgentLoop(
        store=store, model=model, tools=ToolExecutor(store, registry), finalizer=finalize,
    ).run(run["id"], runner_id="test-runner", history=[])
    assert result.status == "finished"
    assert {item["tool_id"] for item in store.actions(run["id"])} == {"query_data", "validate_result"}


def test_agent_cannot_answer_selected_data_question_without_current_query(app):
    store, run = _confirmed_run(app, source_ids=("source-1",), allowed=("query_data",))
    registry = ToolRegistry()
    registry.register(
        ToolSpec("query_data", "query", {"type": "object", "properties": {}}),
        lambda _args: {"output_refs": ["result-1"], "completeness": "complete"},
    )
    model = ScriptedModelAdapter([
        ModelResponse("scripted_test", "fixture", "历史里没找到数据。", (), "stop", None, {"total_tokens": 1}),
        ModelResponse(
            "scripted_test", "fixture", "", (ModelToolCall("q1", "query_data", {}),),
            "tool_calls", None, {"total_tokens": 1},
        ),
        ModelResponse("scripted_test", "fixture", "当前数据已查询。", (), "stop", None, {"total_tokens": 1}),
    ])

    result = AgentLoop(
        store=store, model=model, tools=ToolExecutor(store, registry),
        finalizer=lambda *_args: {"published": True, "quality_status": "passed", "publication_id": "publication-1"},
    ).run(run["id"], runner_id="test-runner", history=[])

    assert result.status == "finished"
    assert result.answer == "当前数据已查询。"
    assert [item["tool_id"] for item in store.actions(run["id"])] == ["query_data"]


def test_query_data_may_narrow_to_one_source_inside_run_scope(app, source):
    from backend.services.advanced_agent import build_executor

    store, run = _confirmed_run(
        app, source_ids=(source["id"], "other-selected-source"), allowed=("query_data",),
    )
    context = store.acquire_lease(run["id"], "source-subset-test")
    decision = store.record_decision(run["id"], ModelResponse(
        "scripted_test", "fixture", "", (), "tool_calls", None, {"total_tokens": 0},
    ))
    executor = build_executor(app.extensions["meridian_db"], store.get_run(run["id"]))
    with app.app_context():
        queried = executor.execute(
            context=context, decision_id=decision["id"], call_id="query",
            tool_id="query_data", arguments={
                "source_ids": [source["id"]], "sql": "SELECT COUNT(*) AS rows FROM data",
            },
        )
    assert queried.result.status.value == "SUCCEEDED"
    assert queried.value["source_ids"] == [source["id"]]

    with app.app_context():
        blocked = executor.execute(
            context=context, decision_id=decision["id"], call_id="outside",
            tool_id="query_data", arguments={
                "source_ids": ["outside-run-scope"], "sql": "SELECT 1 AS value",
            },
        )
    assert blocked.result.status.value == "FAILED"
    assert blocked.result.error_code == "permission_denied"


def test_publication_gate_replays_numeric_claims_against_result_cells(app, source):
    from backend.services.datasets import execute_query

    store, run = _confirmed_run(app, source_ids=(source["id"],))
    with app.app_context():
        result = execute_query(
            [source["id"]], "SELECT SUM(sales) AS total_sales FROM data", "default",
            actor_id="local-default",
        )
        evidence = [
            {
                "tool": "query_data", "status": "SUCCEEDED", "refs": [result["id"]],
                "completeness": "complete", "validation_status": "not_evaluated",
            },
            {
                "tool": "validate_result", "status": "SUCCEEDED", "refs": [result["id"]],
                "completeness": "complete", "validation_status": "PASS",
            },
        ]
        blocked = ResultService(store.db).finalize(run["id"], "销售额合计为 999 元。", evidence)
        assert blocked["published"] is False
        replay = next(
            item for item in blocked["validation"]["items"]
            if item["rule_id"] == "numeric_claim_replay"
        )
        assert replay["status"] == "FAIL"

        published = ResultService(store.db).finalize(run["id"], "销售额合计为 745 元。", evidence)
        assert published["published"] is True
        claims = ResultService(store.db).claims(run["id"], workspace_id="default")
        assert claims[-1]["payload"]["numeric_replay"] == "PASS"
        assert claims[-1]["payload"]["attribution_status"] == "value_only"
        assert claims[-1]["payload"]["evidence_cells"][0]["column"] == "total_sales"
        response = app.test_client().get(
            f"/api/analyses/{run['id']}/evidence/claims/{claims[-1]['id']}/cells/0"
        )
        assert response.status_code == 200
        replayed = response.get_json()["item"]
        assert float(replayed["value"]) == 745
        assert float(replayed["row"]["total_sales"]) == 745
        assert response.get_json()["item"]["result_id"] == result["id"]
        assert app.test_client().get(
            f"/api/analyses/{run['id']}/evidence/claims/{claims[-1]['id']}/cells/1"
        ).status_code == 404
        if len(claims) > 1:
            assert app.test_client().get(
                f"/api/analyses/{run['id']}/evidence/claims/{claims[0]['id']}/cells/0"
            ).status_code == 404


def test_numeric_claim_must_match_the_named_result_row(app, source):
    from backend.services.datasets import execute_query

    store, run = _confirmed_run(app, source_ids=(source["id"],))
    with app.app_context():
        result = execute_query(
            [source["id"]],
            "SELECT region, SUM(sales) AS total_sales FROM data GROUP BY region ORDER BY region",
            "default", actor_id="local-default",
        )
        evidence = [
            {"tool": "query_data", "status": "SUCCEEDED", "refs": [result["id"]], "completeness": "complete"},
            {"tool": "validate_result", "status": "SUCCEEDED", "refs": [result["id"]], "validation_status": "PASS", "completeness": "complete"},
        ]
        wrong_row = ResultService(store.db).finalize(run["id"], "North 的销售额为 305。", evidence)
        assert wrong_row["published"] is False

        right_row = ResultService(store.db).finalize(run["id"], "North 的销售额为 440。", evidence)
        assert right_row["published"] is True


def test_publication_gate_accepts_recovered_auxiliary_failure_and_contract_numbers(app, source):
    from backend.services.datasets import execute_query

    store, run = _confirmed_run(
        app, source_ids=(source["id"],),
        objective="按 60 分及格线分析 2024-2025-1 学期数据",
    )
    with app.app_context():
        result = execute_query(
            [source["id"]],
            "SELECT SUM(sales) AS total_sales FROM data WHERE sales <= 200", "default",
            actor_id="local-default",
        )
        dataset_ref_id = store.db.new_id("dref")
        with store.db.transaction() as connection:
            connection.execute(
                """INSERT INTO dataset_refs(id,workspace_id,run_id,payload,created_at,updated_at)
                   VALUES(?,?,?,?,datetime('now'),datetime('now'))""",
                (
                    dataset_ref_id, "default", run["id"],
                    json.dumps({"location": {"query_result_id": result["id"]}}),
                ),
            )
        evidence = [
            {
                "tool": "query_data", "status": "SUCCEEDED",
                "refs": [dataset_ref_id, result["id"]],
                "completeness": "complete", "validation_status": "not_evaluated",
            },
            {
                "tool": "select_chart", "status": "SUCCEEDED", "refs": [result["id"]],
                "completeness": "unknown", "validation_status": "not_evaluated",
            },
            {
                "tool": "query_data", "status": "FAILED", "refs": [],
                "completeness": "unknown", "validation_status": "not_evaluated",
            },
            {
                "tool": "validate_result", "status": "SUCCEEDED", "refs": [result["id"]],
                "completeness": "complete", "validation_status": "PASS",
            },
        ]
        published = ResultService(store.db).finalize(
            run["id"],
            "1. 2024-2025-1 学期采用 60 分及格线和 200 元查询上限，销售额证据值为 745 元。",
            evidence,
        )
        assert published["published"] is True
        checks = {item["rule_id"]: item for item in published["validation"]["items"]}
        assert checks["tool_failures"]["status"] == "PASS"
        assert checks["result_completeness"]["status"] == "PASS"
        assert checks["numeric_claim_replay"]["status"] == "PASS"


def test_tool_executor_enforces_effective_tools_and_budget(app):
    budget = RunStore.default_budget() | {"tool_calls": 1}
    store, run = _confirmed_run(app, allowed=("allowed",), budget=budget)
    registry = ToolRegistry()
    registry.register(ToolSpec("allowed", "allowed", {"type": "object", "properties": {}}), lambda _args: {"ok": True})
    registry.register(ToolSpec("hidden", "hidden", {"type": "object", "properties": {}}), lambda _args: {"ok": True})
    context = store.acquire_lease(run["id"], "test-runner")
    executor = ToolExecutor(store, registry)
    decision = store.record_decision(run["id"], ModelResponse(
        "scripted_test", "fixture", "", (), "tool_calls", None, {"total_tokens": 0},
    ))
    with pytest.raises(PermissionError):
        executor.execute(context=context, decision_id=decision["id"], call_id="hidden", tool_id="hidden", arguments={})
    executor.execute(context=context, decision_id=decision["id"], call_id="one", tool_id="allowed", arguments={})
    with pytest.raises(RuntimeError, match="预算"):
        executor.execute(context=context, decision_id=decision["id"], call_id="two", tool_id="allowed", arguments={})


def test_stale_lease_is_fenced(app):
    store, run = _confirmed_run(app, allowed=("read",))
    first = store.acquire_lease(run["id"], "runner-one")
    store.release_lease(run["id"], "runner-one", first.lease_epoch)
    store.acquire_lease(run["id"], "runner-two")
    with pytest.raises(PermissionError, match="租约"):
        store.begin_action(
            run["id"], "decision", "logical", "read", {}, lease_epoch=first.lease_epoch,
        )


def test_model_protocol_adapters_preserve_tool_calls_and_usage():
    chunks = [SimpleNamespace(
        usage=None, choices=[SimpleNamespace(finish_reason=None, delta=SimpleNamespace(
            content="依据", refusal=None, tool_calls=[SimpleNamespace(
                index=0, id="call-1", function=SimpleNamespace(name="query", arguments='{"sql":'),
            )],
        ))],
    ), SimpleNamespace(
        usage={"prompt_tokens": 3, "completion_tokens": 2, "total_tokens": 5},
        choices=[SimpleNamespace(finish_reason="tool_calls", delta=SimpleNamespace(
            content=None, refusal=None, tool_calls=[SimpleNamespace(
                index=0, id=None, function=SimpleNamespace(name=None, arguments='"SELECT 1"}'),
            )],
        ))],
    )]
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **_kwargs: iter(chunks))))
    chat = ChatCompletionsAdapter(client, {"model": "test", "temperature": 0})
    response = chat.complete([], [], max_output_tokens=50)
    assert response.tool_calls[0].arguments == {"sql": "SELECT 1"}
    assert response.usage["total_tokens"] == 5

    native = SimpleNamespace(
        status="completed", incomplete_details=None,
        usage={"input_tokens": 4, "output_tokens": 2, "total_tokens": 6},
        output=[SimpleNamespace(type="function_call", call_id="r1", name="query", arguments='{"sql":"SELECT 1"}')],
    )
    responses_client = SimpleNamespace(responses=SimpleNamespace(create=lambda **_kwargs: native))
    response = ResponsesAdapter(responses_client, {"model": "test", "temperature": 0}).complete(
        [], [], max_output_tokens=50,
    )
    assert response.protocol == "openai_responses"
    assert response.tool_calls[0].name == "query"


def test_context_builder_keeps_only_complete_tool_groups_and_respects_budget():
    builder = ContextBuilder(context_window=3000, max_output_tokens=1000)
    incomplete = [
        {"role": "user", "content": "old" * 2000},
        {"role": "assistant", "content": "", "tool_calls": [{"id": "missing"}]},
        {"role": "tool", "tool_call_id": "other", "content": "unpaired"},
        {"role": "user", "content": "latest"},
    ]
    messages = builder.build(
        system="policy", contract={"objective": "test"}, plan=None, history=incomplete,
        evidence_summary=[], skills=[], remaining_budget={"tool_calls": 2},
    )
    assert messages[0]["role"] == "system"
    assert messages[-1]["content"] == "latest"
    assert not any(message.get("tool_call_id") == "other" for message in messages)
    assert not any(message.get("content") == "old" * 2000 for message in messages)


def test_context_builder_compacts_repeated_evidence_previews():
    builder = ContextBuilder(context_window=8000, max_output_tokens=1000)
    messages = builder.build(
        system="policy", contract={"objective": "test"}, plan=None, history=[],
        evidence_summary=[{
            "tool": "query_data", "status": "SUCCEEDED", "refs": [f"result-{index}"],
            "completeness": "complete", "validation_status": "PASS", "preview": "x" * 5000,
        } for index in range(30)],
        skills=[], remaining_budget={"model_tokens": 10000},
    )
    governed = messages[0]["content"]
    assert governed.count("result-") == 24
    assert len(governed) < 40_000


def test_agent_enters_wrap_up_with_evidence_before_budget_is_exhausted():
    evidence = [{"tool": "query_data", "status": "SUCCEEDED", "refs": ["result-1"]}]
    assert not _should_wrap_up({"model_tokens": 100_000}, {"model_tokens": 54_999}, evidence)
    assert _should_wrap_up({"model_tokens": 100_000}, {"model_tokens": 55_000}, evidence)
    assert not _should_wrap_up({"model_tokens": 100_000}, {"model_tokens": 90_000}, [])


def test_agent_tool_messages_are_bounded():
    rendered = _tool_message({"catalog": "x" * 50_000})
    assert len(rendered) < 12_100
    assert "tool result truncated" in rendered


def test_quota_limited_agent_reports_daily_budget_reason(app):
    store, run = _confirmed_run(
        app, allowed=(), budget=RunStore.default_budget() | {"model_tokens": 5_000},
    )
    model = ScriptedModelAdapter([ModelResponse(
        "scripted_test", "fixture", "不会发布", (), "stop", None,
        {"total_tokens": 5_001},
    )])
    result = AgentLoop(
        store=store, model=model, tools=ToolExecutor(store, ToolRegistry()),
        finalizer=lambda *_args: {"published": False},
        model_budget_stop_reason="daily_model_budget_exceeded",
    ).run(run["id"], runner_id="test-runner", history=[])
    assert result.stop_reason == "daily_model_budget_exceeded"


@pytest.mark.parametrize("payload,match", [
    (None, "必须是对象"),
    ({"coverage": "all", "dimensions": ["x"], "deliverables": ["summary"]}, "目标"),
    ({"objective": "x", "dimensions": ["x"], "deliverables": ["summary"]}, "覆盖范围"),
    ({"objective": "x", "coverage": "all", "deliverables": ["summary"]}, "查看维度"),
    ({"objective": "x", "coverage": "all", "dimensions": ["x"]}, "交付形式"),
])
def test_task_contract_rejects_incomplete_sections(payload, match):
    with pytest.raises(ValueError, match=match):
        TaskContract.from_payload(payload)
    with pytest.raises(ValueError, match="字符串"):
        TaskContract.from_payload({
            "objective": "x", "coverage": "all", "dimensions": {"bad": True},
            "deliverables": ["summary"],
        })
    with pytest.raises(ValueError, match="最多选择 100"):
        TaskContract.from_payload({
            "objective": "x", "coverage": "all", "dimensions": ["x"],
            "deliverables": ["summary"], "source_ids": [str(index) for index in range(101)],
        })


def test_model_adapters_reject_malformed_protocol_and_support_fallbacks():
    calls = []

    def create(**kwargs):
        calls.append(kwargs)
        if len(calls) == 1:
            raise TypeError("stream_options unsupported")
        return iter([{"choices": []}])

    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    response = ChatCompletionsAdapter(client, {
        "model": "test", "enable_thinking": True, "reasoning_effort": "high",
    }).complete([], [{"type": "function", "function": {"name": "query"}}], max_output_tokens=20)
    assert response.finish_reason == "stop"
    assert "stream_options" in calls[0] and "stream_options" not in calls[1]
    assert calls[0]["reasoning_effort"] == "high"

    malformed = [{"choices": [{"finish_reason": "tool_calls", "delta": {
        "tool_calls": [{"index": 0, "id": "bad", "function": {"name": "query", "arguments": "[1]"}}],
    }}]}]
    bad_client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **_kwargs: iter(malformed))))
    with pytest.raises(ModelProtocolError, match="JSON 对象"):
        ChatCompletionsAdapter(bad_client, {"model": "test"}).complete([], [], max_output_tokens=20)

    missing_name = [{"choices": [{"finish_reason": "tool_calls", "delta": {
        "tool_calls": [{"index": 0, "id": "bad", "function": {"arguments": "{}"}}],
    }}]}]
    nameless = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(
        create=lambda **_kwargs: iter(missing_name),
    )))
    with pytest.raises(ModelProtocolError, match="没有工具名"):
        ChatCompletionsAdapter(nameless, {"model": "test"}).complete([], [], max_output_tokens=20)

    output = SimpleNamespace(
        status="incomplete", incomplete_details={"reason": "max_output_tokens"},
        usage={"input_tokens": 1, "output_tokens": 2}, output=[SimpleNamespace(
            type="message", content=[
                {"type": "output_text", "text": "visible"}, {"type": "refusal", "text": "limited"},
            ],
        )],
    )
    responses = ResponsesAdapter(
        SimpleNamespace(responses=SimpleNamespace(create=lambda **_kwargs: output)), {"model": "test"},
    ).complete([], [], max_output_tokens=20)
    assert responses.content == "visible" and responses.refusal == "limited"
    assert responses.finish_reason == "length"

    with pytest.raises(InterruptedError):
        ResponsesAdapter(SimpleNamespace(), {"model": "test"}).complete(
            [], [], max_output_tokens=20, should_cancel=lambda: True,
        )
    scripted = ScriptedModelAdapter([])
    with pytest.raises(InterruptedError):
        scripted.complete([], [], max_output_tokens=20, should_cancel=lambda: True)
    assert isinstance(build_model_adapter(SimpleNamespace(), {"model": "x", "protocol": "responses"}), ResponsesAdapter)
    assert isinstance(build_model_adapter(SimpleNamespace(), {"model": "x", "protocol": "chat"}), ChatCompletionsAdapter)
    with pytest.raises(ValueError, match="模型协议"):
        build_model_adapter(SimpleNamespace(), {"model": "x", "protocol": "unknown"})


def test_tool_schema_validation_and_status_mapping(app):
    store, run = _confirmed_run(
        app, allowed=("accepted", "approval", "unknown", "failed", "tuple", "timeout"),
        budget=RunStore.default_budget() | {"tool_calls": 20},
    )
    registry = ToolRegistry()
    schema = {"type": "object", "properties": {}, "required": []}
    registry.register(ToolSpec("accepted", "accepted", schema), lambda _args: {"status": "ACCEPTED", "job_id": "job-1"})
    registry.register(ToolSpec("approval", "approval", schema), lambda _args: {"status": "WAITING_APPROVAL"})
    registry.register(ToolSpec("unknown", "unknown", schema), lambda _args: {"status": "UNKNOWN"})
    registry.register(ToolSpec("failed", "failed", schema), lambda _args: {"ok": False, "error_code": "checked"})
    registry.register(ToolSpec("tuple", "tuple", schema), lambda _args: (
        {"result_id": "result-1", "schema_ref": "schema-1", "preview": {"rows": 1}},
        [("custom", {"ok": True})],
    ))
    registry.register(ToolSpec("timeout", "timeout", schema), lambda _args: (_ for _ in ()).throw(TimeoutError("late")))
    with pytest.raises(ValueError, match="重复"):
        registry.register(ToolSpec("accepted", "duplicate", schema), lambda _args: {})
    with pytest.raises(ValueError, match="未知工具"):
        registry.get("missing")

    context = store.acquire_lease(run["id"], "status-test")
    decision = store.record_decision(run["id"], ModelResponse(
        "scripted_test", "fixture", "", (), "tool_calls", None, {"total_tokens": 0},
    ))
    executor = ToolExecutor(store, registry)
    statuses = {
        name: executor.execute(
            context=context, decision_id=decision["id"], call_id=name, tool_id=name, arguments={},
        ).result.status.value
        for name in ("accepted", "approval", "unknown", "failed", "tuple", "timeout")
    }
    assert statuses == {
        "accepted": "ACCEPTED", "approval": "WAITING_APPROVAL", "unknown": "UNKNOWN",
        "failed": "FAILED", "tuple": "SUCCEEDED", "timeout": "FAILED",
    }
    assert executor.execute(
        context=context, decision_id=decision["id"], call_id="tuple-2", tool_id="tuple", arguments={},
    ).events == (("custom", {"ok": True}),)
    assert executor.effective_tools(context, skill_tools={"tuple"}, child_tools={"tuple", "accepted"}) == {"tuple"}

    with pytest.raises(ValueError, match="必须是对象"):
        _validate([], schema)
    typed = {"required": ["name"], "properties": {
        "name": {"type": "string"}, "items": {"type": "array"}, "meta": {"type": "object"},
        "active": {"type": "boolean"}, "count": {"type": "integer"}, "mode": {"enum": ["safe"]},
    }}
    for value, match in [
        ({}, "缺少字段"), ({"name": 1}, "name"),
        ({"name": "x", "items": {}}, "items"), ({"name": "x", "meta": []}, "meta"),
        ({"name": "x", "active": 1}, "active"), ({"name": "x", "count": True}, "count"),
        ({"name": "x", "mode": "unsafe"}, "允许值"),
    ]:
        with pytest.raises(ValueError, match=match):
            _validate(value, typed)


def test_remote_dataset_ref_never_implicitly_becomes_dataframe(app):
    database = app.extensions["meridian_db"]
    ref = DatasetRef(
        ref_id="remote-ref", kind="remote_objects", source_refs=("warehouse-source",),
        engine_id="livy", location={"output_uri": "s3://results/run/"}, snapshot_set={},
        source_time=None, schema_ref=None, grain=None, query_id="batch-1", query_hash="hash",
        contract_version=1, policy_version="policy", computation_state="complete",
        result_completeness="complete", accuracy="exact", requested_scope={}, actual_scope={},
        sample_metadata={}, row_count=10_000_000, encoded_bytes=1_000_000_000,
        preview_ref=None, provenance_ref="remote_batch:1", retention_until=None,
        owner_id="local-default", acl={"workspace_id": "default", "actor_ids": ["local-default"]},
    )
    DatasetRefStore(database).put(ref, workspace_id="default", run_id=None)
    stored = DatasetRefStore(database).get("remote-ref", workspace_id="default")
    assert stored and stored.kind == "remote_objects"
    assert stored.location["output_uri"].startswith("s3://")


def test_retired_host_mutation_tools_are_physically_absent(app):
    from backend.services.agent_tools import AgentToolContext, tool_schemas

    with app.app_context():
        names = {item["function"]["name"] for item in tool_schemas(
            AgentToolContext(app.extensions["meridian_db"], "default", "missing", []),
        )}
    retired = {
        "workspace_write_file", "workspace_edit_file", "workspace_delete_file", "workspace_bash",
        "configure_hooks", "create_feishu_bitable", "append_feishu_bitable_records",
        "drawio_display", "drawio_edit", "drawio_get",
    }
    assert not (names & retired)
    assert retired.isdisjoint(FORMAL_AGENT_TOOLS)


def test_trino_query_id_cannot_masquerade_as_durable_large_result(app):
    database = app.extensions["meridian_db"]
    query = database.put("warehouse_queries", {
        "id": "trino-large", "workspace_id": "default", "run_id": None,
        "engine_id": "trino", "status": "finished", "next_uri": None,
        "source_refs": ["source-1"], "columns": [{"name": "value"}],
        "preview": [[1]], "stats": {"outputPositions": 1000}, "sql_hash": "hash",
        "result_mode": "preview",
    }, workspace_id="default")
    adapter = TrinoAdapter(database, "default", TrinoConfig(
        engine_id="trino", endpoint="https://trino.example.test", user="test",
        catalog="warehouse", schema="analytics",
    ))
    with pytest.raises(ValueError, match="稳定位置"):
        adapter.result_ref(
            query["id"], owner_id="local-default", contract_version=1, policy_version="policy",
        )
    database.patch("warehouse_queries", query["id"], {
        "result_mode": "materialize", "materialized_location": {
            "catalog": "warehouse", "schema": "meridian_results", "table": "result_1",
        }, "update_count": 1000,
    }, workspace_id="default")
    ref = adapter.result_ref(
        query["id"], owner_id="local-default", contract_version=1, policy_version="policy",
    )
    assert ref.kind == "remote_table"
    assert ref.row_count == 1000


def test_http_cursor_pagination_records_truthful_completion(app, monkeypatch):
    import json

    from backend.services.datasets import _http_pagination, _load_http_dataset

    payloads = [
        {"items": [{"id": 1}, {"id": 2}], "meta": {"next": "cursor-2"}},
        {"items": [{"id": 3}], "meta": {"next": None}},
    ]

    class Response:
        def __init__(self, payload):
            self._payload = payload
            self.content = json.dumps(payload).encode()
            self.text = self.content.decode()
            self.headers = {"Content-Type": "application/json"}

        def json(self):
            return self._payload

        def raise_for_status(self):
            return None

    calls = []

    def request(_method, _url, **kwargs):
        calls.append(kwargs.get("params") or {})
        return Response(payloads[len(calls) - 1])

    monkeypatch.setattr("backend.services.datasets.safe_http_request", request)
    with app.app_context():
        config = _http_pagination({"pagination": {
            "mode": "cursor", "cursor_param": "after",
            "next_cursor_path": "meta.next", "max_pages": 10,
        }})
        frame, meta = _load_http_dataset(
            "https://api.example.test/items", {}, "items", config,
        )
    assert frame["id"].tolist() == [1, 2, 3]
    assert calls == [{}, {"after": "cursor-2"}]
    assert meta == {
        "mode": "cursor", "pages_fetched": 2, "records_fetched": 3,
        "complete": True, "limit_reason": None,
    }


def test_formal_python_analysis_uses_only_bounded_sandbox(app, source, monkeypatch):
    from backend.services.advanced_agent import build_executor

    database = app.extensions["meridian_db"]
    store, run = _confirmed_run(
        app, source_ids=(source["id"],), allowed=("query_data", "run_analysis"),
    )
    context = store.acquire_lease(run["id"], "sandbox-test")
    decision = store.record_decision(run["id"], ModelResponse(
        "scripted_test", "fixture", "", (), "tool_calls", None, {"total_tokens": 0},
    ))
    executor = build_executor(database, store.get_run(run["id"]))
    with app.app_context():
        queried = executor.execute(
            context=context, decision_id=decision["id"], call_id="query",
            tool_id="query_data", arguments={"sql": "SELECT * FROM data"},
        )
    input_ref = queried.value["dataset_ref_id"]

    monkeypatch.setattr(
        "backend.services.advanced_agent.execute_tool",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("host analysis path called")),
    )
    monkeypatch.setenv("MERIDIAN_SANDBOX_PROXY_URL", "http://127.0.0.1:8090")
    monkeypatch.setenv("MERIDIAN_SANDBOX_PROXY_TOKEN", "test-sandbox-token-that-is-long-enough")

    def fake_execute(self, spec, *, input_dir, run_id, should_cancel=None):
        assert spec["method"] == "describe"
        assert not should_cancel()
        output_dir = self.output_root / run_id
        output_dir.mkdir()
        target = output_dir / "result.parquet"
        import pandas as pd

        pd.DataFrame({"metric": ["rows"], "value": [6]}).to_parquet(target, index=False)
        return {
            "status": "SUCCEEDED", "output_dir": str(output_dir),
            "files": [{
                "path": "result.parquet", "bytes": target.stat().st_size,
                "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
            }],
            "metrics": {"input_rows": 6, "output_rows": 1, "method": "describe"},
        }

    monkeypatch.setattr("backend.services.data_plane.sandbox_client.SandboxClient.execute", fake_execute)
    with app.app_context():
        analyzed = executor.execute(
            context=context, decision_id=decision["id"], call_id="analysis",
            tool_id="run_analysis", arguments={"dataset_ref_id": input_ref, "method": "describe"},
        )
    assert analyzed.result.status.value == "SUCCEEDED"
    result_ref = DatasetRefStore(database).get(analyzed.value["dataset_ref_id"], workspace_id="default")
    assert result_ref and result_ref.result_completeness == "complete"
    assert analyzed.value["provenance_ref"].startswith("sandbox:")


def test_formal_python_analysis_fails_closed_without_sandbox(app, source, monkeypatch):
    from backend.services.advanced_agent import build_executor

    database = app.extensions["meridian_db"]
    store, run = _confirmed_run(
        app, source_ids=(source["id"],), allowed=("query_data", "run_analysis"),
    )
    context = store.acquire_lease(run["id"], "sandbox-unavailable-test")
    decision = store.record_decision(run["id"], ModelResponse(
        "scripted_test", "fixture", "", (), "tool_calls", None, {"total_tokens": 0},
    ))
    executor = build_executor(database, store.get_run(run["id"]))
    with app.app_context():
        queried = executor.execute(
            context=context, decision_id=decision["id"], call_id="query",
            tool_id="query_data", arguments={"sql": "SELECT * FROM data"},
        )

    def unavailable(*_args, **_kwargs):
        raise SandboxUnavailable("隔离容器不可用")

    monkeypatch.setenv("MERIDIAN_SANDBOX_PROXY_URL", "http://127.0.0.1:8090")
    monkeypatch.setenv("MERIDIAN_SANDBOX_PROXY_TOKEN", "test-sandbox-token-that-is-long-enough")
    monkeypatch.setattr("backend.services.data_plane.sandbox_client.SandboxClient.execute", unavailable)
    with app.app_context():
        analyzed = executor.execute(
            context=context, decision_id=decision["id"], call_id="analysis",
            tool_id="run_analysis", arguments={
                "dataset_ref_id": queried.value["dataset_ref_id"],
                "code": "result = df.describe().reset_index()",
            },
        )
    assert analyzed.result.status.value == "FAILED"
    assert "隔离容器不可用" in analyzed.value["error"]
    assert not database.list("analysis_runs", workspace_id="default")
