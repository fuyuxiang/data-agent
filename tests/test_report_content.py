from __future__ import annotations

import copy
import io
import json
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from backend.agent.contracts import ModelResponse, ModelToolCall, TaskContract
from backend.agent.loop import AgentLoop
from backend.agent.store import RunStore
from backend.services.advanced_agent import available_formal_tools, build_executor
from backend.services.results.content import checked_report, current_result_ids, normalize_report, submit_report
from backend.services.results.manifests import ResultService
from backend.skills.loader import load_builtin_skills, read_reference
from backend.skills.registry import SkillRegistry


def formal_run(app, client):
    response = client.post('/api/sources/upload', data={
        'file': (io.BytesIO(b'group,amount\nA,10\nB,20\nA,0\nB,-5\n'), 'measurements.csv'),
    })
    source = response.get_json()['items'][0]
    database = app.extensions['meridian_db']
    session = database.put('sessions', {'id': database.new_id('ses'), 'workspace_id': 'default',
                                      'owner_id': 'local-default'}, workspace_id='default')
    store = RunStore(database)
    with app.app_context():
        tools = available_formal_tools(database, 'default', session['id'], [source['id']])
    run, _ = store.create_run(workspace_id='default', session_id=session['id'], actor_id='local-default',
                              source_scope=[source['id']], allowed_tool_ids=tools)
    store.add_contract(run['id'], TaskContract.from_payload({
        'objective': '比较各组的数值并生成报告', 'coverage': '已选来源全部记录',
        'dimensions': ['group'], 'source_scope': [source['id']], 'deliverables': ['summary', 'report'],
    }), expected_version=0, confirmed_by='local-default')
    return store, run


class ReportModel:
    protocol, model = 'fixture', 'fixture'

    def __init__(self, store, run):
        self.store, self.run, self.stage = store, run, 0
        self.report = None

    def complete(self, _messages, schemas, **_kwargs):
        available = {item['function']['name'] for item in schemas}
        assert 'submit_analysis_report' in available
        if self.stage == 0:
            calls = [ModelToolCall('total', 'query_data', {
                'sql': 'SELECT SUM(amount) AS total, 0 AS zero, CAST(NULL AS DOUBLE) AS missing, -5 AS negative FROM measurements',
            }),
                     ModelToolCall('groups', 'query_data', {
                         'sql': 'SELECT "group", SUM(amount) AS amount FROM measurements GROUP BY "group" ORDER BY "group"',
                     })]
        elif self.stage == 1:
            ids = current_result_ids(self.store.db, self.run)
            refs = [next(ref for ref in ids if 'total' in self.store.db.get('query_results', ref)['columns']),
                    next(ref for ref in ids if 'group' in self.store.db.get('query_results', ref)['columns'])]
            self.report = {
                'title': '分组数值分析报告',
                'metrics': [{'name': '合计', 'cell': {'result_id': refs[0], 'row': 0, 'column': 'total'},
                             'definition': '已选记录的数值合计', 'formula': 'SUM(amount)', 'unit': '单位'}],
                'findings': [{'text': 'B 组的 amount 为 15。', 'type': 'fact',
                              'confidence': 'high', 'confidence_reason': '直接核对完整分组查询结果',
                              'evidence_refs': [refs[1]]},
                             {'text': '分组差异可能与记录构成有关。', 'type': 'inference',
                              'confidence': 'low', 'confidence_reason': '缺少独立影响因素数据',
                              'evidence_refs': [refs[1]]}],
                'recommendations': [{'text': '复核负值记录的业务含义', 'owner': '数据负责人',
                                     'scope': '当前数据范围', 'priority': 'medium', 'evidence_refs': [refs[1]]}],
                'glossary': [{'name': '分组数值', 'result_id': refs[1], 'column': 'amount',
                              'definition': '同组记录的数值合计', 'formula': 'SUM(amount)', 'unit': '单位'}],
                'tables': [{'result_id': refs[0], 'title': '整体统计'}, {'result_id': refs[1], 'title': '分组明细'}],
                'assumptions': ['负值按原始数值保留'], 'limitations': ['业务准确性尚未独立评估'],
            }
            calls = [ModelToolCall('profile', 'profile_data', {'result_id': refs[1]}),
                     ModelToolCall('chart', 'generate_chart', {
                'result_id': refs[1], 'type': 'grouped_bar', 'x': 'group', 'y': 'amount', 'title': '分组数值对比',
            }), ModelToolCall('report', 'submit_analysis_report', self.report)]
        else:
            self.stage += 1
            return ModelResponse('fixture', 'fixture', '合计为 25。B 组的 amount 为 15。', (), 'stop', None,
                                 {'total_tokens': 20})
        self.stage += 1
        return ModelResponse('fixture', 'fixture', '', tuple(calls), 'tool_calls', None, {'total_tokens': 20})


def run_report(app, client):
    store, run = formal_run(app, client)
    model = ReportModel(store, run)
    with app.app_context():
        result = AgentLoop(store=store, model=model, tools=build_executor(store.db, run),
                           finalizer=ResultService(store.db).finalize).run(run['id'], runner_id='report-fixture', history=[])
    assert result.outcome == 'complete', result
    return store, run, model.report


def test_report_contract_runs_through_real_queries_and_all_export_formats(app, client, monkeypatch):
    store, run, report = run_report(app, client)
    response = client.get(f"/api/analyses/{run['id']}/results")
    payload = response.get_json()['manifest']['payload']
    assert payload['kpis'][0]['value'] == 25
    assert payload['kpis'][0]['unit'] == '单位'
    assert len(payload['tables']) == 2
    assert payload['charts'][0]['type'] == 'grouped_bar'
    assert payload['report']['findings'][1]['type'] == 'inference'
    assert payload['data_quality']['accuracy'] == '未评估'
    assert payload['data_quality']['checks'][0]['rows'] == 2
    second_id = report['tables'][1]['result_id']
    page = client.get(f"/api/analyses/{run['id']}/details?result_id={second_id}&cursor=1&limit=1")
    assert page.get_json()['items'] == [{'group': 'B', 'amount': 15.0}]
    assert client.get(f"/api/analyses/{run['id']}/details?result_id=unrelated").status_code == 403

    # Exporting must consume the publication, without querying a source or invoking the model again.
    monkeypatch.setattr('backend.services.datasets.execute_query', lambda *_a, **_k: pytest.fail('重新取数'))
    artifacts = client.post(f"/api/analyses/{run['id']}/artifacts", json={
        'kinds': ['report_docx', 'report_html', 'report_pptx', 'data_xlsx'],
    })
    assert artifacts.status_code == 201, artifacts.get_json()
    assert all(item['title'].startswith('分组数值分析报告') for item in artifacts.get_json()['items'])
    assert all('manifest_' not in item['filename'] for item in artifacts.get_json()['items'])
    paths = {item['kind']: store.db.get('artifacts', item['id'], workspace_id='default')['path']
             for item in artifacts.get_json()['items']}
    word = Document(paths['report_docx'])
    word_text = '\n'.join(paragraph.text for paragraph in word.paragraphs)
    slides = Presentation(paths['report_pptx'])
    slides_text = '\n'.join(shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame)
    html_text = open(paths['report_html'], encoding='utf-8').read()
    for text in (word_text, slides_text, html_text):
        assert 'B 组的 amount 为 15' in text
        assert '业务准确性' in text
        assert '复核负值记录的业务含义' in text
        assert 'SUM(amount)' in text
    assert 'data:image/png;base64,' in html_text
    assert '<td>缺失</td>' in html_text and '<td>0</td>' in html_text
    assert '<script' not in html_text and 'cdn.' not in html_text
    book = load_workbook(paths['data_xlsx'], read_only=True)
    assert book['验证数据']['A2'].value == 25
    assert book['验证数据 2']['B3'].value == 15
    assert '指标与字段说明' in book.sheetnames
    assert '行动建议' in book.sheetnames
    assert client.post(f"/api/analyses/{run['id']}/artifacts", json={'kinds': ['report_docx']}).get_json()['items'][0]['id'] == next(
        item['id'] for item in artifacts.get_json()['items'] if item['kind'] == 'report_docx')


@pytest.mark.parametrize('mutation, message', [
    (lambda report: report['metrics'][0].update(value=999), '不支持的字段'),
    (lambda report: report['metrics'][0]['cell'].update(row=-1), '行号'),
    (lambda report: report['metrics'][0]['cell'].update(column='unknown'), '字段不存在'),
    (lambda report: report['findings'][0].update(text='amount 为 999。'), '数字无法'),
    (lambda report: report['findings'][0].update(evidence_refs=[]), '证据引用'),
    (lambda report: report['findings'][1].update(confidence='high'), '高置信度'),
    (lambda report: report.update(quality={'score': 'A'}), '不支持的字段'),
    (lambda report: report['metrics'][0].update(definition='<script>alert(1)</script>'), 'HTML'),
])
def test_report_rejects_invented_values_and_unsupported_quality_claims(app, client, mutation, message):
    store, run, report = run_report(app, client)
    report = copy.deepcopy(report)
    mutation(report)
    with app.app_context(), pytest.raises(ValueError, match=message):
        normalize_report(store.db, run, report, current_result_ids(store.db, run))


def test_report_cannot_borrow_another_run_or_use_truncated_evidence(app, client):
    store, run, report = run_report(app, client)
    _, other, _ = run_report(app, client)
    with app.app_context():
        with pytest.raises(PermissionError, match='当前分析'):
            normalize_report(store.db, other, report, current_result_ids(store.db, other))
        ref = report['metrics'][0]['cell']['result_id']
        store.db.patch('query_results', ref, {'completeness': 'system_truncated'}, workspace_id='default')
        with pytest.raises(ValueError, match='截断'):
            normalize_report(store.db, run, report, current_result_ids(store.db, run))


def test_reference_sections_survive_clone_export_and_import(client, app):
    builtin = load_builtin_skills()['data-analysis']
    section = read_reference(builtin, 'methods.md', section='诊断与对比', limit=3)
    assert section['next_offset'] == 3
    assert '结构与异常' not in section['content']
    with pytest.raises(ValueError):
        read_reference(builtin, '../secrets.txt')
    clone = client.post('/api/skills/data-analysis/clone').get_json()['item']
    exported = client.get(f"/api/skills/{clone['id']}/export")
    archive = zipfile.ZipFile(io.BytesIO(exported.data))
    assert any(name.endswith('/references/methods.md') for name in archive.namelist())
    assert client.delete(f"/api/skills/{clone['id']}").status_code == 200
    imported = client.post('/api/skills/import', data={'file': (io.BytesIO(exported.data), 'skill.zip')})
    assert imported.status_code == 201, imported.get_json()
    with app.app_context():
        definition = SkillRegistry(app.extensions['meridian_db'], 'default').get(imported.get_json()['item']['id'])
    assert read_reference(definition, 'methods.md', section='结构与异常')['content'].startswith('## 结构与异常')


def test_published_report_is_immutable_and_changed_evidence_cannot_be_reused(app, client):
    from backend.services.query_results import read_result_frame

    store, run, report = run_report(app, client)
    with app.app_context():
        with pytest.raises(ValueError, match='已结束'):
            submit_report(store.db, run, report)
        result = store.db.get('query_results', report['metrics'][0]['cell']['result_id'], workspace_id='default')
        frame = read_result_frame(result)
        frame.loc[0, 'total'] = 26
        frame.to_parquet(result['path'], index=False)
        with pytest.raises(ValueError, match='证据已变化'):
            checked_report(store.db, run, current_result_ids(store.db, run))


def test_metric_cells_preserve_zero_missing_negative_and_plain_text_formulas(app, client):
    store, run, report = run_report(app, client)
    ref = report['metrics'][0]['cell']['result_id']
    with app.app_context():
        content = normalize_report(store.db, run, {'metrics': [
            {'name': column, 'cell': {'result_id': ref, 'row': 0, 'column': column},
             'formula': 'CASE WHEN x < 10 AND y > 2 THEN x END'}
            for column in ['zero', 'missing', 'negative']
        ]}, current_result_ids(store.db, run))
    assert [item['value'] for item in content['metrics']] == [0, None, -5]
    json.dumps(content, allow_nan=False)


def test_failed_report_revision_cannot_silently_publish_an_older_draft(app, client):
    class RevisionModel(ReportModel):
        def complete(self, messages, schemas, **kwargs):
            if self.stage == 2:
                self.stage += 1
                invalid = copy.deepcopy(self.report)
                invalid['metrics'][0]['value'] = 999
                return ModelResponse('fixture', 'fixture', '', (
                    ModelToolCall('invalid-revision', 'submit_analysis_report', invalid),
                ), 'tool_calls', None, {'total_tokens': 20})
            return super().complete(messages, schemas, **kwargs)

    store, run = formal_run(app, client)
    with app.app_context():
        result = AgentLoop(store=store, model=RevisionModel(store, run), tools=build_executor(store.db, run),
                           finalizer=ResultService(store.db).finalize).run(run['id'], runner_id='revision', history=[])
    assert result.publication_id is None
    assert result.quality_status == 'failed'
    assert client.get(f"/api/analyses/{run['id']}/results").get_json()['status'] == 'not_published'


def test_governed_metric_definition_is_frozen_at_query_time(app, client):
    from backend.services.results.content import field_definition

    store, run = formal_run(app, client)
    model_response = client.post('/api/semantic/models', json={
        'name': 'measurement-model', 'source_id': run['source_scope'][0], 'table': 'measurements',
        'grain': '每行一条观测记录', 'dimensions': [],
        'measures': [{'name': 'amount_sum', 'column': 'amount', 'aggregation': 'sum'}],
    })
    assert model_response.status_code == 201, model_response.get_json()
    metric_response = client.post('/api/semantic/metrics', json={
        'name': 'total_amount', 'label': '数值合计', 'description': '已选观测数值的合计',
        'model_id': model_response.get_json()['item']['id'], 'measure': 'amount_sum',
        'status': 'approved', 'unit': '单位',
    })
    assert metric_response.status_code == 201, metric_response.get_json()
    metric = metric_response.get_json()['item']
    query = client.post('/api/semantic/query', json={'metric': metric['id']})
    assert query.status_code == 200, query.get_json()
    result = query.get_json()['result']
    assert client.patch(f"/api/semantic/metrics/{metric['id']}", json={
        'label': '更新后的名称', 'description': '更新后的定义', 'unit': '新单位',
    }).status_code == 200
    with app.app_context():
        entry = field_definition(store.db, 'default', result, 'total_amount')
        content = normalize_report(store.db, run, {'metrics': [{
            'name': '数值合计', 'cell': {'result_id': result['id'], 'row': 0, 'column': 'total_amount'},
            'definition': '模型猜测的定义', 'unit': '错误单位',
        }]}, [result['id']])
    assert entry['definition'] == '已选观测数值的合计'
    assert entry['unit'] == '单位'
    assert 'SUM' in entry['formula']
    assert entry['metric_version'] == 1
    assert content['metrics'][0]['definition_status'] == 'governed'
    assert content['metrics'][0]['unit'] == '单位'


def report_followup(app, parent, *, source_scope=None, time_range=None, session_id=None):
    database = app.extensions['meridian_db']
    store = RunStore(database)
    scope = parent['source_scope'] if source_scope is None else source_scope
    with app.app_context():
        tools = available_formal_tools(database, 'default', session_id or parent['session_id'], scope)
    run, _ = store.create_run(workspace_id='default', session_id=session_id or parent['session_id'],
                              actor_id='local-default', source_scope=scope, allowed_tool_ids=tools, skill_id='report')
    store.add_contract(run['id'], TaskContract.from_payload({
        'objective': '把刚才的分析整理成报告', 'coverage': '沿用已选来源', 'dimensions': ['整体'],
        'source_scope': scope, 'deliverables': ['report'], 'time_range': time_range or {},
    }), expected_version=0, confirmed_by='local-default')
    return store, run


def test_report_only_followup_reuses_authorized_evidence_without_querying(app, client, monkeypatch):
    from backend.services.advanced_agent import _prepare_run_skills

    store, parent, report = run_report(app, client)
    store, run = report_followup(app, parent)

    class ReuseModel:
        protocol, model, stage = 'fixture', 'fixture', 0

        def complete(self, _messages, schemas, **_kwargs):
            names = {item['function']['name'] for item in schemas}
            assert 'reuse_analysis_result' in names
            if self.stage == 0:
                calls = [ModelToolCall('reuse', 'reuse_analysis_result', {})]
            elif self.stage == 1:
                reused = next(action['result']['value'] for action in store.actions(run['id'])
                              if action['tool_id'] == 'reuse_analysis_result')
                assert reused['report_input'] == report
                calls = [ModelToolCall('compose', 'submit_analysis_report', reused['report_input'])]
            else:
                return ModelResponse('fixture', 'fixture', '合计为 25。B 组的 amount 为 15。', (),
                                     'stop', None, {'total_tokens': 20})
            self.stage += 1
            return ModelResponse('fixture', 'fixture', '', tuple(calls), 'tool_calls', None, {'total_tokens': 20})

    monkeypatch.setattr('backend.services.datasets.execute_query', lambda *_a, **_k: pytest.fail('报告重新取数'))
    with app.app_context():
        skills = _prepare_run_skills(store.db, run, {'skill_ids': ['report']}, run['allowed_tool_ids'])
        result = AgentLoop(store=store, model=ReuseModel(), tools=build_executor(store.db, run),
                           finalizer=ResultService(store.db).finalize).run(run['id'], runner_id='reuse', history=[], skills=skills)
    assert result.outcome == 'complete', result
    payload = client.get(f"/api/analyses/{run['id']}/results").get_json()['manifest']['payload']
    assert payload['kpis'][0]['value'] == 25
    assert len(payload['tables']) == 2
    assert any('未重新取数' in value for value in payload['limitations'])
    assert {action['tool_id'] for action in store.actions(run['id'])} == {'reuse_analysis_result', 'submit_analysis_report'}


def test_new_report_skill_can_analyze_selected_data_when_no_prior_publication_exists(app, client):
    from backend.services.advanced_agent import _prepare_run_skills

    store, run = formal_run(app, client)
    with app.app_context():
        skills = _prepare_run_skills(store.db, run, {'skill_ids': ['report']}, run['allowed_tool_ids'])
        result = AgentLoop(store=store, model=ReportModel(store, run), tools=build_executor(store.db, run),
                           finalizer=ResultService(store.db).finalize).run(run['id'], runner_id='fresh-report',
                                                                        history=[], skills=skills)
    assert result.outcome == 'complete', result
    assert client.get(f"/api/analyses/{run['id']}/results").get_json()['manifest']['payload']['kpis'][0]['value'] == 25


def test_result_reuse_respects_session_scope_and_new_time_range(app, client):
    from backend.services.advanced_agent import _reuse_analysis_result

    store, parent, _report = run_report(app, client)
    _, narrowed = report_followup(app, parent, source_scope=[])
    _, changed_time = report_followup(app, parent, time_range={'start': '2026-01-01', 'end': '2026-02-01'})
    session = store.db.put('sessions', {'id': store.db.new_id('ses'), 'workspace_id': 'default',
                                      'owner_id': 'local-default'}, workspace_id='default')
    _, other_session = report_followup(app, parent, session_id=session['id'])
    with app.app_context():
        with pytest.raises(PermissionError, match='数据范围'):
            _reuse_analysis_result(store.db, narrowed)
        with pytest.raises(ValueError, match='时间范围'):
            _reuse_analysis_result(store.db, changed_time)
        with pytest.raises(ValueError, match='当前会话'):
            _reuse_analysis_result(store.db, other_session)


def test_reference_payload_rejects_paths_and_is_not_injected_whole_into_model(app):
    from backend.skills.executor import SkillExecutor
    from backend.skills.context import SkillContext
    from backend.skills.models import SkillError, skill_from_payload

    with pytest.raises(SkillError, match='文件名'):
        skill_from_payload({'id': 'safe-skill', 'name': '通用', 'description': '通用', 'instruction': '分析',
                            'reference_documents': {'../private.md': 'secret'}})
    builtin = load_builtin_skills()['report']
    block = SkillExecutor(SkillContext(workspace_id='default', actor_id='local-default', run_id='fixture',
                                      session_id='fixture')).prepare([builtin])[0].model_blocks[0]
    assert block['references'][0]['name'] == 'content-contract.md'
    assert 'reference_documents' not in block
    assert len(json.dumps(block, ensure_ascii=False)) < 5000
